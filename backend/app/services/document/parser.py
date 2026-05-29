import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def parse_pdf(file_path: str) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(file_path)
    text_parts = []
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            text_parts.append(f"[Page {i + 1}]\n{page_text}")
    return "\n\n".join(text_parts)


def parse_docx(file_path: str) -> str:
    from docx import Document

    doc = Document(file_path)
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            text_parts.append(row_text)
    return "\n\n".join(text_parts)


def parse_xlsx(file_path: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True)
    text_parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        text_parts.append(f"[Sheet: {sheet_name}]")
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
            if row_text.strip(" |"):
                text_parts.append(row_text)
    wb.close()
    return "\n\n".join(text_parts)


def parse_image_ocr(file_path: str) -> str:
    try:
        import pytesseract
        from PIL import Image

        image = Image.open(file_path)
        text = pytesseract.image_to_string(image)
        return text
    except Exception as e:
        logger.warning(f"OCR failed for {file_path}: {e}")
        return ""


def parse_document(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    parsers = {
        ".pdf": parse_pdf,
        ".docx": parse_docx,
        ".xlsx": parse_xlsx,
        ".xls": parse_xlsx,
        ".png": parse_image_ocr,
        ".jpg": parse_image_ocr,
        ".jpeg": parse_image_ocr,
        ".tiff": parse_image_ocr,
        ".tif": parse_image_ocr,
    }
    parser = parsers.get(suffix)
    if parser is None:
        if suffix == ".txt":
            return path.read_text(encoding="utf-8", errors="ignore")
        raise ValueError(f"Unsupported file format: {suffix}")
    return parser(file_path)


def chunk_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> list[dict]:
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_text(text)
    result = []
    for i, chunk in enumerate(chunks):
        page_num = None
        if "[Page " in chunk:
            try:
                page_marker = chunk.split("[Page ")[1].split("]")[0]
                page_num = int(page_marker)
            except (IndexError, ValueError):
                pass
        result.append({"index": i, "content": chunk, "page_number": page_num})
    return result
